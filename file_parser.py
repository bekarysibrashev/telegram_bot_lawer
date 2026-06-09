import io
import logging

logger = logging.getLogger(__name__)
MAX_CHARS = 12000


async def extract_text_from_bytes(file_bytes: bytes, filename: str) -> str:
    filename_lower = filename.lower()
    if filename_lower.endswith(".pdf"):
        return _extract_pdf(file_bytes, filename)
    elif filename_lower.endswith(".docx"):
        return _extract_docx(file_bytes, filename)
    return ""


def _extract_pdf(file_bytes: bytes, filename: str) -> str:
    try:
        import fitz
        doc = fitz.open(stream=file_bytes, filetype="pdf")
        pages_text = [page.get_text() for page in doc if page.get_text().strip()]
        doc.close()
        return _truncate("\n".join(pages_text).strip(), filename)
    except Exception as e:
        logger.error(f"PDF extract error: {e}")
        return f"[Ошибка чтения PDF: {e}]"


def _extract_docx(file_bytes: bytes, filename: str) -> str:
    try:
        from docx import Document
        doc = Document(io.BytesIO(file_bytes))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(c.text.strip() for c in row.cells if c.text.strip())
                if row_text:
                    parts.append(row_text)
        return _truncate("\n".join(parts).strip(), filename)
    except Exception as e:
        logger.error(f"DOCX extract error: {e}")
        return f"[Ошибка чтения DOCX: {e}]"


def _truncate(text: str, filename: str) -> str:
    if not text:
        return f"[Файл '{filename}' не содержит текста или защищён паролем]"
    if len(text) > MAX_CHARS:
        text = text[:MAX_CHARS] + f"\n\n[... документ обрезан до {MAX_CHARS} символов ...]"
    return text
